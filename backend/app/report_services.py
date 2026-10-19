from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from . import models, schemas


ICTU_TIMEZONE = "Asia/Ho_Chi_Minh"
ICTU_BLUE = "00529B"
ICTU_NAVY = "003B70"
ICTU_ORANGE = "F5A800"
ICTU_LIGHT_BLUE = "EAF4FB"
ICTU_BORDER = "B8CEDF"


def _logo_path() -> Path | None:
    candidates = [
        Path(__file__).resolve().parent / "assets" / "ICTU.png",
        Path(__file__).resolve().parents[2] / "frontend" / "public" / "assets" / "ICTU.png",
    ]
    return next((candidate for candidate in candidates if candidate.exists()), None)


def _local_time(value: datetime) -> datetime:
    return value.astimezone(ZoneInfo(ICTU_TIMEZONE))


def _period_label(report: schemas.ReportOverviewOut) -> str:
    start = _local_time(report.date_from).strftime("%d/%m/%Y")
    end = _local_time(report.date_to).strftime("%d/%m/%Y")
    return f"Kỳ báo cáo: {start} - {end}"


def _duration_minutes(start: datetime, end: datetime) -> int:
    return max(0, int((end - start).total_seconds() // 60))


def report_overview(
    db: Session,
    *,
    date_from: datetime,
    date_to: datetime,
    room_id: int | None = None,
    floor: int | None = None,
    building: str | None = None,
) -> schemas.ReportOverviewOut:
    if date_to < date_from:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Khoảng báo cáo không hợp lệ")

    stmt = (
        select(models.Meeting)
        .options(
            selectinload(models.Meeting.booking).selectinload(models.RoomBooking.room),
            selectinload(models.Meeting.participants),
        )
        .where(models.Meeting.start_time >= date_from, models.Meeting.start_time <= date_to)
        .order_by(models.Meeting.start_time)
    )
    meetings = list(db.scalars(stmt).all())

    normalized_building = building.strip().lower() if building else None
    if room_id is not None or floor is not None or normalized_building:
        meetings = [
            meeting
            for meeting in meetings
            if meeting.booking
            and (room_id is None or meeting.booking.room_id == room_id)
            and (floor is None or meeting.booking.room.floor == floor)
            and (
                normalized_building is None
                or meeting.booking.room.building.strip().lower() == normalized_building
            )
        ]

    now = datetime.now(timezone.utc)
    total = len(meetings)
    cancelled = sum(item.status == models.MeetingStatus.CANCELLED for item in meetings)
    scheduled = sum(
        item.status == models.MeetingStatus.SCHEDULED and item.end_time >= now
        for item in meetings
    )
    completed = sum(
        item.status == models.MeetingStatus.SCHEDULED and item.end_time < now
        for item in meetings
    )
    bookings = [item.booking for item in meetings if item.booking is not None]

    room_buckets: dict[int, dict] = {}
    for meeting in meetings:
        if meeting.booking is None:
            continue
        room = meeting.booking.room
        bucket = room_buckets.setdefault(
            room.id,
            {
                "room_id": room.id,
                "room_name": room.name,
                "building": room.building,
                "floor": room.floor,
                "booking_count": 0,
                "booked_minutes": 0,
                "scheduled_count": 0,
                "completed_count": 0,
                "cancelled_count": 0,
            },
        )
        bucket["booking_count"] += 1
        bucket["booked_minutes"] += _duration_minutes(meeting.start_time, meeting.end_time)
        if meeting.status == models.MeetingStatus.CANCELLED:
            bucket["cancelled_count"] += 1
        elif meeting.end_time < now:
            bucket["completed_count"] += 1
        else:
            bucket["scheduled_count"] += 1

    room_usage = [schemas.RoomUsageRow(**value) for value in room_buckets.values()]
    room_usage.sort(key=lambda row: (-row.booking_count, -row.booked_minutes, row.room_name))

    daily: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    reasons: Counter[str] = Counter()
    organizers: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for meeting in meetings:
        day = meeting.start_time.astimezone(ZoneInfo(ICTU_TIMEZONE)).date().isoformat()
        daily[day][0] += 1
        organizers[meeting.organizer_email][0] += 1
        if meeting.status == models.MeetingStatus.CANCELLED:
            daily[day][1] += 1
            organizers[meeting.organizer_email][1] += 1
            reasons[(meeting.cancellation_reason or "Không ghi lý do").strip() or "Không ghi lý do"] += 1

    trend = [
        schemas.CancellationTrendRow(
            date=day,
            total_count=counts[0],
            cancelled_count=counts[1],
            cancellation_rate=round(counts[1] * 100 / counts[0], 2) if counts[0] else 0,
        )
        for day, counts in sorted(daily.items())
    ]
    organizer_rows = [
        schemas.OrganizerCancellationRow(
            organizer_email=email,
            total_count=counts[0],
            cancelled_count=counts[1],
            cancellation_rate=round(counts[1] * 100 / counts[0], 2) if counts[0] else 0,
        )
        for email, counts in sorted(organizers.items(), key=lambda item: (-item[1][1], item[0]))
    ]

    return schemas.ReportOverviewOut(
        generated_at=now,
        date_from=date_from,
        date_to=date_to,
        summary=schemas.ReportSummary(
            total_meetings=total,
            scheduled_meetings=scheduled,
            completed_meetings=completed,
            cancelled_meetings=cancelled,
            cancellation_rate=round(cancelled * 100 / total, 2) if total else 0,
            total_bookings=len(bookings),
            total_booking_minutes=sum(
                _duration_minutes(item.start_time, item.end_time) for item in bookings
            ),
        ),
        room_usage=room_usage,
        cancellation_trend=trend,
        cancellation_reasons=[
            schemas.CancellationReasonRow(reason=reason, count=count)
            for reason, count in reasons.most_common()
        ],
        organizer_cancellations=organizer_rows,
        top_room=room_usage[0] if room_usage else None,
    )


def _excel_brand_header(sheet, title: str, report: schemas.ReportOverviewOut, last_column: int) -> None:
    last_letter = get_column_letter(last_column)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_column)
    sheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_column)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=last_column)
    sheet["B1"] = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG"
    sheet["B2"] = "HỆ THỐNG QUẢN LÝ LỊCH HỌP ICTU"
    sheet["A3"] = title
    sheet["A4"] = (
        f"{_period_label(report)} | Xuất lúc "
        f"{_local_time(report.generated_at).strftime('%H:%M %d/%m/%Y')} | Múi giờ: UTC+7"
    )
    sheet["B1"].font = Font(name="Arial", size=12, bold=True, color=ICTU_NAVY)
    sheet["B2"].font = Font(name="Arial", size=10, italic=True, color="4C6478")
    sheet["B1"].alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    sheet["B2"].alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    sheet["A3"].font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    sheet["A3"].fill = PatternFill("solid", fgColor=ICTU_BLUE)
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A4"].font = Font(name="Arial", size=9, italic=True, color="4C6478")
    sheet["A4"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 20
    sheet.row_dimensions[3].height = 30
    sheet.row_dimensions[4].height = 22
    sheet.print_title_rows = "1:6"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = "Hệ thống quản lý lịch họp ICTU"
    sheet.oddFooter.right.text = "Trang &P / &N"
    sheet.oddFooter.center.size = 8
    sheet.oddFooter.right.size = 8
    sheet.auto_filter.ref = None

    logo = _logo_path()
    if logo is not None:
        image = ExcelImage(str(logo))
        image.width = 55
        image.height = 55
        sheet.add_image(image, "A1")
    else:
        sheet["A1"] = "ICTU"
        sheet["A1"].font = Font(name="Arial", size=14, bold=True, color=ICTU_BLUE)
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")

    sheet.print_area = f"A1:{last_letter}{max(sheet.max_row, 6)}"


def _style_excel_table(sheet, header_row: int, last_row: int, last_column: int) -> None:
    header_fill = PatternFill("solid", fgColor=ICTU_BLUE)
    stripe_fill = PatternFill("solid", fgColor="F3F8FC")
    border = Border(bottom=Side(style="thin", color=ICTU_BORDER))
    for cell in sheet[header_row][:last_column]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 28
    for row_index in range(header_row + 1, last_row + 1):
        for column_index in range(1, last_column + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.font = Font(name="Arial", size=10, color="243B53")
            cell.alignment = Alignment(
                horizontal="left" if column_index == 1 else "center",
                vertical="center",
                wrap_text=column_index == 1,
            )
            cell.border = border
            if row_index % 2 == 0:
                cell.fill = stripe_fill
        sheet.row_dimensions[row_index].height = 23


def _set_excel_widths(sheet, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def export_excel(report: schemas.ReportOverviewOut) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "Meeting Management ICTU"
    workbook.properties.title = "Báo cáo sử dụng phòng họp ICTU"
    workbook.properties.subject = _period_label(report)
    workbook.calculation.fullCalcOnLoad = True

    summary_sheet = workbook.active
    summary_sheet.title = "Tổng quan"
    _excel_brand_header(summary_sheet, "BÁO CÁO SỬ DỤNG PHÒNG HỌP", report, 7)
    metric_labels = [
        "Tổng cuộc họp", "Đang/sắp diễn ra", "Đã hoàn thành", "Đã hủy",
        "Tỷ lệ hủy", "Lượt đặt phòng", "Giờ sử dụng",
    ]
    metric_values = [
        report.summary.total_meetings,
        report.summary.scheduled_meetings,
        report.summary.completed_meetings,
        report.summary.cancelled_meetings,
        report.summary.cancellation_rate / 100,
        report.summary.total_bookings,
        report.summary.total_booking_minutes / 60,
    ]
    for column, (label, value) in enumerate(zip(metric_labels, metric_values, strict=True), start=1):
        label_cell = summary_sheet.cell(row=6, column=column, value=label)
        value_cell = summary_sheet.cell(row=7, column=column, value=value)
        label_cell.fill = PatternFill("solid", fgColor=ICTU_LIGHT_BLUE)
        label_cell.font = Font(name="Arial", size=9, bold=True, color=ICTU_NAVY)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell.font = Font(name="Arial", size=16, bold=True, color=ICTU_BLUE)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.border = Border(bottom=Side(style="medium", color=ICTU_ORANGE))
    summary_sheet["E7"].number_format = "0.0%"
    summary_sheet["G7"].number_format = "0.0"
    summary_sheet.row_dimensions[6].height = 32
    summary_sheet.row_dimensions[7].height = 34
    summary_sheet.merge_cells("A9:G9")
    summary_sheet["A9"] = (
        f"Phòng được sử dụng nhiều nhất: {report.top_room.room_name} "
        f"({report.top_room.booking_count} lượt đặt)"
        if report.top_room else "Chưa có dữ liệu sử dụng phòng trong kỳ báo cáo."
    )
    summary_sheet["A9"].font = Font(name="Arial", size=10, italic=True, color="4C6478")
    summary_sheet["A9"].alignment = Alignment(horizontal="left")
    summary_sheet.freeze_panes = "A6"
    _set_excel_widths(summary_sheet, [19, 19, 19, 16, 16, 18, 18])
    summary_sheet.print_area = "A1:G10"

    room_sheet = workbook.create_sheet("Sử dụng phòng")
    _excel_brand_header(room_sheet, "CHI TIẾT SỬ DỤNG PHÒNG HỌP", report, 8)
    room_headers = ["Phòng", "Khu nhà", "Tầng", "Lượt đặt", "Giờ sử dụng", "Sắp tới", "Hoàn thành", "Đã hủy"]
    room_sheet.append([])
    room_sheet.append(room_headers)
    room_header_row = 6
    if report.room_usage:
        for row in report.room_usage:
            room_sheet.append([
                row.room_name, row.building, row.floor, row.booking_count,
                row.booked_minutes / 60, row.scheduled_count, row.completed_count,
                row.cancelled_count,
            ])
    else:
        room_sheet.append(["Không có dữ liệu", "", "", 0, 0, 0, 0, 0])
    _style_excel_table(room_sheet, room_header_row, room_sheet.max_row, 8)
    for cell in room_sheet["E"][room_header_row:]:
        cell.number_format = "0.0"
    room_sheet.freeze_panes = "A7"
    room_sheet.auto_filter.ref = f"A{room_header_row}:H{room_sheet.max_row}"
    _set_excel_widths(room_sheet, [25, 20, 10, 14, 17, 14, 16, 13])
    room_sheet.print_area = f"A1:H{room_sheet.max_row}"

    cancel_sheet = workbook.create_sheet("Thống kê hủy")
    _excel_brand_header(cancel_sheet, "THỐNG KÊ CUỘC HỌP BỊ HỦY", report, 4)
    current_row = 6
    cancel_sheet.cell(current_row, 1, "Xu hướng hủy theo ngày")
    cancel_sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    cancel_sheet.cell(current_row, 1).font = Font(name="Arial", size=11, bold=True, color=ICTU_NAVY)
    current_row += 1
    trend_header = current_row
    for column, value in enumerate(["Ngày", "Tổng cuộc họp", "Đã hủy", "Tỷ lệ hủy"], start=1):
        cancel_sheet.cell(current_row, column, value)
    for row in report.cancellation_trend:
        current_row += 1
        cancel_sheet.append([datetime.fromisoformat(row.date).date(), row.total_count, row.cancelled_count, row.cancellation_rate / 100])
    if current_row == trend_header:
        current_row += 1
        cancel_sheet.append(["Không có dữ liệu", 0, 0, 0])
    _style_excel_table(cancel_sheet, trend_header, current_row, 4)
    for cell in cancel_sheet["A"][trend_header:]:
        if not isinstance(cell.value, str):
            cell.number_format = "dd/mm/yyyy"
    for cell in cancel_sheet["D"][trend_header:]:
        cell.number_format = "0.0%"

    current_row += 2
    cancel_sheet.cell(current_row, 1, "Lý do hủy")
    cancel_sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    cancel_sheet.cell(current_row, 1).font = Font(name="Arial", size=11, bold=True, color=ICTU_NAVY)
    current_row += 1
    reason_header = current_row
    cancel_sheet.cell(current_row, 1, "Lý do")
    cancel_sheet.cell(current_row, 2, "Số lần")
    for row in report.cancellation_reasons:
        current_row += 1
        cancel_sheet.cell(current_row, 1, row.reason)
        cancel_sheet.cell(current_row, 2, row.count)
    if current_row == reason_header:
        current_row += 1
        cancel_sheet.cell(current_row, 1, "Không có dữ liệu")
        cancel_sheet.cell(current_row, 2, 0)
    _style_excel_table(cancel_sheet, reason_header, current_row, 2)

    current_row += 2
    cancel_sheet.cell(current_row, 1, "Hủy theo người tổ chức")
    cancel_sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    cancel_sheet.cell(current_row, 1).font = Font(name="Arial", size=11, bold=True, color=ICTU_NAVY)
    current_row += 1
    organizer_header = current_row
    for column, value in enumerate(["Người tổ chức", "Tổng cuộc họp", "Đã hủy", "Tỷ lệ hủy"], start=1):
        cancel_sheet.cell(current_row, column, value)
    for row in report.organizer_cancellations:
        current_row += 1
        cancel_sheet.append([row.organizer_email, row.total_count, row.cancelled_count, row.cancellation_rate / 100])
    if current_row == organizer_header:
        current_row += 1
        cancel_sheet.append(["Không có dữ liệu", 0, 0, 0])
    _style_excel_table(cancel_sheet, organizer_header, current_row, 4)
    for cell in cancel_sheet["D"][organizer_header:]:
        cell.number_format = "0.0%"
    cancel_sheet.freeze_panes = "A7"
    _set_excel_widths(cancel_sheet, [42, 20, 16, 18])
    cancel_sheet.print_area = f"A1:D{current_row}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_font() -> str:
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            if "ICTUUnicode" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ICTUUnicode", str(candidate)))
            return "ICTUUnicode"
    return "Helvetica"


def _pdf_page_decorator(report: schemas.ReportOverviewOut, font: str):
    logo = _logo_path()

    def draw_page(canvas, document) -> None:
        page_width, page_height = landscape(A4)
        canvas.saveState()
        if logo is not None:
            watermark_size = 92 * mm
            if hasattr(canvas, "setFillAlpha"):
                canvas.setFillAlpha(0.055)
            canvas.drawImage(
                str(logo),
                (page_width - watermark_size) / 2,
                (page_height - watermark_size) / 2 - 6 * mm,
                width=watermark_size,
                height=watermark_size,
                preserveAspectRatio=True,
                mask="auto",
            )
            if hasattr(canvas, "setFillAlpha"):
                canvas.setFillAlpha(1)
            canvas.drawImage(
                str(logo), 12 * mm, page_height - 21 * mm,
                width=13 * mm, height=13 * mm, preserveAspectRatio=True, mask="auto",
            )
        canvas.setFillColor(colors.HexColor(f"#{ICTU_NAVY}"))
        canvas.setFont(font, 9)
        canvas.drawString(29 * mm, page_height - 13 * mm, "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG")
        canvas.setFillColor(colors.HexColor("#567086"))
        canvas.setFont(font, 7.5)
        canvas.drawString(29 * mm, page_height - 17.5 * mm, "Hệ thống quản lý lịch họp ICTU")
        canvas.setStrokeColor(colors.HexColor(f"#{ICTU_BLUE}"))
        canvas.setLineWidth(0.8)
        canvas.line(12 * mm, page_height - 23 * mm, page_width - 12 * mm, page_height - 23 * mm)

        canvas.setStrokeColor(colors.HexColor("#C9D8E5"))
        canvas.setLineWidth(0.4)
        canvas.line(12 * mm, 10 * mm, page_width - 12 * mm, 10 * mm)
        canvas.setFillColor(colors.HexColor("#657B8E"))
        canvas.setFont(font, 7)
        canvas.drawString(12 * mm, 6 * mm, _period_label(report))
        canvas.drawRightString(
            page_width - 12 * mm,
            6 * mm,
            f"Trang {document.page} | Meeting Management ICTU",
        )
        canvas.restoreState()

    return draw_page


def _pdf_table(data: list[list], font: str, widths: list[float]) -> Table:
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{ICTU_BLUE}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), font),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 7),
        ("TOPPADDING", (0, 0), (-1, 0), 7),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.HexColor(f"#{ICTU_ORANGE}")),
        ("LINEBELOW", (0, 1), (-1, -1), 0.35, colors.HexColor("#C9D8E5")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F8FC")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
    ]))
    return table


def export_pdf(report: schemas.ReportOverviewOut) -> bytes:
    output = BytesIO()
    font = _pdf_font()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=29 * mm,
        bottomMargin=15 * mm,
        title="Báo cáo sử dụng phòng họp ICTU",
        author="Meeting Management ICTU",
        subject=_period_label(report),
    )
    styles = getSampleStyleSheet()
    styles["Title"].fontName = font
    styles["Normal"].fontName = font
    styles["Heading2"].fontName = font
    styles["Title"].fontSize = 18
    styles["Title"].leading = 22
    styles["Title"].textColor = colors.HexColor(f"#{ICTU_NAVY}")
    styles["Title"].spaceAfter = 3 * mm
    styles["Heading2"].fontSize = 11
    styles["Heading2"].leading = 14
    styles["Heading2"].textColor = colors.HexColor(f"#{ICTU_NAVY}")
    styles["Heading2"].spaceBefore = 3 * mm
    styles["Heading2"].spaceAfter = 2 * mm
    styles.add(ParagraphStyle(
        name="ICTUMeta",
        parent=styles["Normal"],
        fontName=font,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#536B7E"),
        alignment=1,
    ))
    styles.add(ParagraphStyle(
        name="ICTUNote",
        parent=styles["Normal"],
        fontName=font,
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#405A70"),
        backColor=colors.HexColor("#EAF4FB"),
        borderColor=colors.HexColor("#B8CEDF"),
        borderWidth=0.5,
        borderPadding=7,
    ))
    story = [
        Paragraph("BÁO CÁO SỬ DỤNG PHÒNG HỌP", styles["Title"]),
        Paragraph(
            f"{_period_label(report)}<br/>"
            f"Dữ liệu tổng hợp lúc {_local_time(report.generated_at).strftime('%H:%M, ngày %d/%m/%Y')} (UTC+7)",
            styles["ICTUMeta"],
        ),
        Spacer(1, 4 * mm),
    ]
    summary_data = [
        ["Tổng cuộc họp", "Sắp tới", "Hoàn thành", "Đã hủy", "Tỷ lệ hủy", "Lượt đặt", "Giờ sử dụng"],
        [
            report.summary.total_meetings,
            report.summary.scheduled_meetings,
            report.summary.completed_meetings,
            report.summary.cancelled_meetings,
            f"{report.summary.cancellation_rate:.1f}%",
            report.summary.total_bookings,
            f"{report.summary.total_booking_minutes / 60:.1f}",
        ],
    ]
    room_data = [["Phòng", "Khu nhà", "Tầng", "Lượt đặt", "Giờ sử dụng", "Sắp tới", "Hoàn thành", "Đã hủy"]]
    room_data.extend([
        [
            row.room_name, row.building, row.floor, row.booking_count,
            f"{row.booked_minutes / 60:.1f}", row.scheduled_count,
            row.completed_count, row.cancelled_count,
        ]
        for row in report.room_usage
    ])
    if len(room_data) == 1:
        room_data.append(["Không có dữ liệu", "", "", "", "", "", "", ""])
    trend_data = [["Ngày", "Tổng cuộc họp", "Đã hủy", "Tỷ lệ hủy"]]
    trend_data.extend([
        [datetime.fromisoformat(row.date).strftime("%d/%m/%Y"), row.total_count, row.cancelled_count, f"{row.cancellation_rate:.1f}%"]
        for row in report.cancellation_trend
    ])
    if len(trend_data) == 1:
        trend_data.append(["Không có dữ liệu", "", "", ""])
    reason_data = [["Lý do hủy", "Số lần"]]
    reason_data.extend([[row.reason, row.count] for row in report.cancellation_reasons])
    if len(reason_data) == 1:
        reason_data.append(["Không có dữ liệu", ""])
    organizer_data = [["Người tổ chức", "Tổng cuộc họp", "Đã hủy", "Tỷ lệ hủy"]]
    organizer_data.extend([
        [row.organizer_email, row.total_count, row.cancelled_count, f"{row.cancellation_rate:.1f}%"]
        for row in report.organizer_cancellations
    ])
    if len(organizer_data) == 1:
        organizer_data.append(["Không có dữ liệu", "", "", ""])

    story.append(Paragraph("Tổng quan", styles["Heading2"]))
    story.append(_pdf_table(summary_data, font, [36 * mm] * 7))
    story.append(Spacer(1, 4 * mm))
    top_room_text = (
        f"Phòng được sử dụng nhiều nhất: <b>{report.top_room.room_name}</b> "
        f"với {report.top_room.booking_count} lượt đặt và {report.top_room.booked_minutes / 60:.1f} giờ sử dụng."
        if report.top_room else "Chưa có dữ liệu sử dụng phòng trong kỳ báo cáo."
    )
    story.extend([
        Paragraph(top_room_text, styles["ICTUNote"]),
        Paragraph("Chi tiết sử dụng phòng", styles["Heading2"]),
        _pdf_table(room_data, font, [44 * mm, 35 * mm, 17 * mm, 26 * mm, 32 * mm, 28 * mm, 31 * mm, 24 * mm]),
        PageBreak(),
        Spacer(1, 18 * mm),
        Paragraph("Thống kê cuộc họp bị hủy", styles["Title"]),
        Paragraph("Các số liệu dưới đây hỗ trợ theo dõi xu hướng và nguyên nhân hủy lịch họp.", styles["ICTUMeta"]),
        Spacer(1, 3 * mm),
        Paragraph("Xu hướng hủy theo ngày", styles["Heading2"]),
        _pdf_table(trend_data, font, [65 * mm, 50 * mm, 45 * mm, 45 * mm]),
        Spacer(1, 4 * mm),
        KeepTogether([
            Paragraph("Lý do hủy", styles["Heading2"]),
            _pdf_table(reason_data, font, [180 * mm, 35 * mm]),
        ]),
        Spacer(1, 4 * mm),
        KeepTogether([
            Paragraph("Hủy theo người tổ chức", styles["Heading2"]),
            _pdf_table(organizer_data, font, [130 * mm, 40 * mm, 35 * mm, 40 * mm]),
        ]),
    ])

    decorate = _pdf_page_decorator(report, font)
    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return output.getvalue()
